#!/usr/bin/env python3
"""Rebuild the translations workbook: one sheet per landing + one for the thank-you page."""
import re, json, copy
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "Claude outputs" / "Unitty Solutions — Translations (EN_KA_RU).xlsx"
LANDINGS = ["start", "crm", "ai", "banks", "solutions"]
TITLES = {"start": "Start", "crm": "CRM", "ai": "AI", "banks": "Banks", "solutions": "Solutions"}

def dct(path):
    s = path.read_text(encoding="utf-8")
    return json.loads(re.search(r"const I18N = (\{.*?\});\n", s).group(1))

wb = load_workbook(XLSX)
src = wb["Untitled"] if "Untitled" in wb.sheetnames else wb["Start"]
hdr_style = [copy.copy(c._style) for c in src[1][:8]]
row_style = [copy.copy(c._style) for c in src[2][:8]]
widths = {k: v.width for k, v in src.column_dimensions.items()}
prev = {}
for r in src.iter_rows(min_row=2, max_col=8, values_only=True):
    if r[0] and r[2]: prev[(r[0], r[2])] = r
if "Thank-you" in wb.sheetnames:
    for r in wb["Thank-you"].iter_rows(min_row=2, max_col=8, values_only=True):
        if r[2]: prev[("Thank-you", r[2])] = r

def fresh(title, index):
    if title in wb.sheetnames: del wb[title]
    ws = wb.create_sheet(title, index)
    for j, h in enumerate(["Page", "Section", "Key", "English (source)", "Georgian (ქართული)", "Russian (Русский)", "Notes / Context", "Status"]):
        c = ws.cell(row=1, column=j + 1, value=h); c._style = copy.copy(hdr_style[j])
    for k, w in widths.items(): ws.column_dimensions[k].width = w
    ws.freeze_panes = "D2"
    return ws

def write(ws, rows, highlight=set()):
    fill = PatternFill("solid", fgColor="FFF4CC")
    for i, r in enumerate(rows, start=2):
        for j, v in enumerate(r):
            c = ws.cell(row=i, column=j + 1, value=v); c._style = copy.copy(row_style[j])
            if r[2] in highlight and j in (3, 4, 5): c.fill = fill
    ws.auto_filter.ref = f"A1:H{len(rows) + 1}"
    dv = DataValidation(type="list", formula1='"Done,Updated,New,Changed"', allow_blank=True)
    dv.add(f"H2:H{len(rows) + 200}"); ws.add_data_validation(dv)

start = dct(ROOT / "public" / "index.html")
idx = 1
for name in LANDINGS:
    path = ROOT / "public" / ("index.html" if name == "start" else f"{name}/index.html")
    d = dct(path)
    cfgp = ROOT / "landings" / f"{name}.json"
    own = set(json.loads(cfgp.read_text(encoding="utf-8")).get("i18n", {})) if cfgp.exists() else set()
    has_pains = bool(json.loads(cfgp.read_text(encoding="utf-8")).get("pains")) if cfgp.exists() else False
    rows = []
    for k in d["en"]:
        if k.startswith("ty."): continue
        if k.startswith("pain.") and not has_pains: continue   # pain cards only on landings that show them
        p = prev.get(("Landing", k))
        section = p[1] if p else k.split(".")[0].capitalize()
        note = p[6] if p and name == "start" else ("Only on this landing" if k in own else (p[6] if p else ""))
        status = (p[7] if p and name == "start" else ("New" if k in own else (p[7] if p else "New")))
        rows.append([TITLES[name], section, k, d["en"][k], d["ka"][k], d["ru"][k], note, status])
    ws = fresh(TITLES[name], idx); idx += 1
    write(ws, rows, highlight=own if name != "start" else set())
    print(f"+++ {TITLES[name]}: {len(rows)} rows, {len(own)} landing-specific")

ty = dct(ROOT / "public" / "thank-you" / "index.html")
tys = (ROOT / "public" / "thank-you" / "index.html").read_text(encoding="utf-8")
used = set(re.findall(r'data-i18n(?:-html)?="([^"]+)"', tys)) | {"ty.confirm.title", "ty.confirm.title.noname"}
rows = []
for k in ty["en"]:
    if not (k.startswith("ty.") or k in used): continue
    p = prev.get(("Thank-you", k))
    rows.append(["Thank-you (all landings)", p[1] if p else k.split(".")[1].capitalize(), k, ty["en"][k], ty["ka"][k], ty["ru"][k], p[6] if p else "", p[7] if p else "New"])
ws = fresh("Thank-you", idx); write(ws, rows)
print(f"+++ Thank-you: {len(rows)} rows")
if "Untitled" in wb.sheetnames: del wb["Untitled"]

h = wb["How to edit"]
for r in range(1, 30):
    h.cell(row=r, column=1).value = None
lines = [
 "How to edit this file",
 "Each landing has its own sheet: Start (start.unitty.io), CRM, AI, Banks, Solutions. The Thank-you sheet is shared by all landings.",
 "1. Edit text only in columns D (English), E (Georgian), F (Russian).",
 "2. Set Status (column H) to \"Changed\" on every row you edit - only those rows are applied to the site.",
 "3. Do not change column C (Key) - it links the text to its place on the page.",
 "4. Yellow cells = texts that exist only on that landing. Every other row is shared with the Start page until you change it on that sheet.",
 "5. Keep placeholders such as {name} exactly as they are.",
 "6. Send the file back to Claude.",
 "",
 "Status values: Done = live and unchanged · Updated = changed recently · New = added recently · Changed = you edited it, to be applied",
 "\"Removed keys\" sheet: texts that are no longer on the site (kept for reference only).",
]
for i, t in enumerate(lines, start=1): h.cell(row=i, column=1, value=t)
order = ["How to edit"] + [TITLES[n] for n in LANDINGS] + ["Thank-you", "Removed keys"]
wb._sheets = [wb[n] for n in order]
wb.active = 1
wb.save(XLSX)
print("*** saved", XLSX.name, wb.sheetnames)
